"""
Tests for DocumentParser:
Text files, Word documents (.docx), PDFs, and OCR image files.
"""

import io
import pytest
from PIL import Image, ImageDraw
import docx
from pypdf import PdfWriter
from dlp.parser import DocumentParser


class TestDocumentParser:
    def setup_method(self):
        self.parser = DocumentParser()

    def test_parse_plain_text(self):
        sample = b"This is a confidential text file with Aadhaar: 2345 6789 0124"
        doc = self.parser.parse_file(sample, "audit.txt")

        assert doc.filename == "audit.txt"
        assert "2345 6789 0124" in doc.raw_text
        assert doc.ocr_applied is False

    def test_parse_docx(self):
        # Create an in-memory docx
        doc_obj = docx.Document()
        doc_obj.add_paragraph("Customer Record:")
        doc_obj.add_paragraph("Visa Card: 4111 1111 1111 1111")
        table = doc_obj.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "PAN"
        table.rows[0].cells[1].text = "ABCPS1234F"

        buf = io.BytesIO()
        doc_obj.save(buf)
        buf.seek(0)

        parsed = self.parser.parse_file(buf.getvalue(), "accounts.docx")
        assert "accounts.docx" == parsed.filename
        assert "4111 1111 1111 1111" in parsed.raw_text
        assert "ABCPS1234F" in parsed.raw_text
        assert parsed.ocr_applied is False

    def test_parse_image_ocr(self):
        # Render an image with clear text
        img = Image.new("RGB", (600, 150), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        draw.text((30, 50), "CONFIDENTIAL PAN: ABCPS1234F", fill=(0, 0, 0))

        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)

        parsed = self.parser.parse_file(buf.getvalue(), "scan_card.png")
        assert parsed.ocr_applied is True
        assert "ABCPS1234F" in parsed.raw_text.replace(" ", "")

    def test_parse_pdf_native(self):
        # Create a small valid PDF with pypdf
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buf = io.BytesIO()
        writer.write(buf)
        buf.seek(0)

        parsed = self.parser.parse_file(buf.getvalue(), "sample.pdf")
        assert parsed.file_type == "PDF Document"
