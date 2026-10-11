"""
Document and File Ingestion Pipeline for DLP Scanner.
Extracts raw text from:
- Plain text / code / logs / csv (.txt, .log, .csv, .json, .xml, .md, etc.)
- Microsoft Word (.docx)
- Adobe PDF (.pdf) with dual-mode (Native text extraction + Scanned OCR fallback)
- Image files (.png, .jpg, .jpeg, .bmp, .tiff, .webp) via Tesseract OCR
"""

import io
import os
import logging
from dataclasses import dataclass, field
from typing import Optional, Dict, Any, Tuple
from PIL import Image, ImageOps, ImageEnhance
import pytesseract
import pypdf
from pdf2image import convert_from_bytes
import docx

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    filename: str
    file_type: str
    raw_text: str
    ocr_applied: bool
    page_count: int = 1
    metadata: Dict[str, Any] = field(default_factory=dict)


class DocumentParser:
    """
    Parses documents and applies OCR where necessary.
    """

    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp"}
    TEXT_EXTENSIONS = {".txt", ".csv", ".tsv", ".log", ".json", ".xml", ".html", ".htm", ".md", ".yaml", ".yml", ".py", ".sql"}
    DOCX_EXTENSIONS = {".docx"}
    PDF_EXTENSIONS = {".pdf"}

    def __init__(self, tesseract_cmd: Optional[str] = None):
        if tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    def parse_file(
        self,
        file_bytes: bytes,
        filename: str,
        force_ocr: bool = False,
    ) -> ParsedDocument:
        """
        Main entry point to parse a file from bytes.
        """
        ext = os.path.splitext(filename.lower())[1]

        if ext in self.PDF_EXTENSIONS:
            return self._parse_pdf(file_bytes, filename, force_ocr=force_ocr)
        elif ext in self.IMAGE_EXTENSIONS:
            return self._parse_image(file_bytes, filename)
        elif ext in self.DOCX_EXTENSIONS:
            return self._parse_docx(file_bytes, filename)
        else:
            # Fallback to text parsing
            return self._parse_text(file_bytes, filename)

    def _parse_text(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        """
        Decode plain text with multiple encoding fallbacks.
        """
        encodings = ["utf-8", "utf-16", "latin-1", "cp1252", "ascii"]
        text = ""
        used_encoding = "utf-8"

        for enc in encodings:
            try:
                text = file_bytes.decode(enc)
                used_encoding = enc
                break
            except (UnicodeDecodeError, LookupError):
                continue

        # If all decodings failed, force decode with replacement
        if not text and file_bytes:
            text = file_bytes.decode("utf-8", errors="replace")
            used_encoding = "utf-8 (lossy)"

        return ParsedDocument(
            filename=filename,
            file_type="Plain Text / Document",
            raw_text=text,
            ocr_applied=False,
            page_count=1,
            metadata={"encoding": used_encoding, "size_bytes": len(file_bytes)},
        )

    def _parse_docx(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        """
        Extract text from Microsoft Word .docx files.
        """
        try:
            doc_stream = io.BytesIO(file_bytes)
            doc = docx.Document(doc_stream)

            full_text = []

            # Paragraphs
            for p in doc.paragraphs:
                if p.text:
                    full_text.append(p.text)

            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        full_text.append(" | ".join(row_data))

            text = "\n".join(full_text)
            return ParsedDocument(
                filename=filename,
                file_type="Microsoft Word (.docx)",
                raw_text=text,
                ocr_applied=False,
                page_count=1,
                metadata={
                    "paragraphs_count": len(doc.paragraphs),
                    "tables_count": len(doc.tables),
                    "size_bytes": len(file_bytes),
                },
            )
        except Exception as e:
            logger.warning("DOCX extraction error for %s: %s. Falling back to text decode.", filename, e)
            return self._parse_text(file_bytes, filename)

    def _parse_pdf(self, file_bytes: bytes, filename: str, force_ocr: bool = False) -> ParsedDocument:
        """
        Extract text from PDF.
        If PDF is scanned or text layer is scarce, automatically triggers OCR via pdf2image & Tesseract.
        """
        extracted_pages = []
        page_count = 0
        is_scanned = False
        ocr_applied = False

        if not force_ocr:
            try:
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                page_count = len(reader.pages)
                for i, page in enumerate(reader.pages):
                    page_text = page.extract_text() or ""
                    extracted_pages.append(page_text.strip())

                total_chars = sum(len(p) for p in extracted_pages)
                # If very few characters extracted per page on average (< 25 chars), likely scanned
                if page_count > 0 and (total_chars / page_count) < 25:
                    is_scanned = True
            except Exception as e:
                logger.info("PDF reader failed on %s: %s. Attempting OCR fallback.", filename, e)
                is_scanned = True

        # Run OCR if forced or detected as scanned
        if force_ocr or is_scanned:
            try:
                ocr_applied = True
                images = convert_from_bytes(file_bytes, dpi=200)
                page_count = len(images)
                extracted_pages = []

                for page_num, img in enumerate(images, start=1):
                    ocr_text = self._ocr_image_pil(img)
                    extracted_pages.append(f"--- [Page {page_num}] ---\n{ocr_text}")
            except Exception as e:
                logger.error("OCR extraction failed for PDF %s: %s", filename, e)
                # If OCR failed, keep whatever partial text was read from pypdf
                if not extracted_pages:
                    extracted_pages = [f"[Error during PDF extraction: {e}]"]

        full_text = "\n\n".join(extracted_pages)
        return ParsedDocument(
            filename=filename,
            file_type="PDF Document",
            raw_text=full_text,
            ocr_applied=ocr_applied,
            page_count=page_count,
            metadata={
                "is_scanned_pdf": is_scanned,
                "pages_detected": page_count,
                "size_bytes": len(file_bytes),
            },
        )

    def _parse_image(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        """
        Parse image using PIL and Tesseract OCR with image enhancement.
        """
        try:
            image = Image.open(io.BytesIO(file_bytes))
            ocr_text = self._ocr_image_pil(image)

            return ParsedDocument(
                filename=filename,
                file_type="Image (OCR Processed)",
                raw_text=ocr_text,
                ocr_applied=True,
                page_count=1,
                metadata={
                    "dimensions": f"{image.width}x{image.height}",
                    "format": image.format or "Unknown",
                    "mode": image.mode,
                    "size_bytes": len(file_bytes),
                },
            )
        except Exception as e:
            logger.error("Image parsing failed for %s: %s", filename, e)
            return ParsedDocument(
                filename=filename,
                file_type="Image",
                raw_text="",
                ocr_applied=True,
                page_count=1,
                metadata={"error": str(e), "size_bytes": len(file_bytes)},
            )

    def _ocr_image_pil(self, image: Image.Image) -> str:
        """
        Preprocesses image for OCR and runs Tesseract.
        """
        # Convert to RGB if needed
        if image.mode not in ("RGB", "L"):
            image = image.convert("RGB")

        # Scale up small images for better OCR accuracy
        if image.width < 800 or image.height < 600:
            factor = max(1.5, min(3.0, 1200 / max(image.width, 1)))
            new_size = (int(image.width * factor), int(image.height * factor))
            image = image.resize(new_size, Image.Resampling.LANCZOS)

        # Grayscale
        gray = image.convert("L")

        # Enhance contrast
        enhancer = ImageEnhance.Contrast(gray)
        enhanced = enhancer.enhance(1.8)

        # First pass with standard page segmentation
        try:
            text = pytesseract.image_to_string(enhanced, config="--psm 3")
        except Exception:
            text = pytesseract.image_to_string(image)

        # If sparse, try psm 6 (assume uniform block of text)
        if len(text.strip()) < 10:
            try:
                alt_text = pytesseract.image_to_string(enhanced, config="--psm 6")
                if len(alt_text.strip()) > len(text.strip()):
                    text = alt_text
            except Exception:
                pass

        return text.strip()
